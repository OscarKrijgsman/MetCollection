import re
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
import xlrd

from .normalize import (
    catalog_norm,
    clean,
    content_hash,
    extract_barcode,
    match_format,
    parse_quantity,
    parse_year,
)

TEXT_FIELDS = ["label", "country", "catalog_no", "matrix", "made_in", "description"]


def load_workbook(path):
    """Sheet name (trimmed) -> list of rows, each a list of raw cell values."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(str(path))
        return {s.name.strip(): [s.row_values(r) for r in range(s.nrows)] for s in book.sheets()}
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {ws.title.strip(): [list(r) for r in ws.iter_rows(values_only=True)] for ws in book.worksheets}


def column_letter(index):
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


class Result:
    def __init__(self):
        self.items = []
        self.concerts = []
        self.sheet_counts = {}
        self.errors = []
        self.warnings = []


def read_items(sheet_rows, spec, eras, seen_ids, result):
    sheet = spec["sheet"]
    columns = spec["columns"]
    header = [clean(v) or "" for v in sheet_rows[0]] if sheet_rows else []
    for i, (expected, _) in enumerate(columns):
        got = header[i] if i < len(header) else None
        if got is None or got.lower() != expected.lower():
            result.errors.append(
                f"{sheet}: column {column_letter(i)} header is {got!r}, expected {expected!r}. "
                "If the layout changed on purpose, update config/sheets.yaml."
            )
            return
    trailing = list(enumerate(header))[len(columns):]
    id_re = re.compile(rf"^{re.escape(spec['prefix'])}-\d{{4,}}$")
    unknown_eras = defaultdict(list)
    count = 0
    total = None

    for row_no, row in enumerate(sheet_rows[1:], start=2):
        raw = {field: row[i] if i < len(row) else None for i, (_, field) in enumerate(columns)}
        item_id = clean(raw["id"])
        title = clean(raw["title"])
        quantity = parse_quantity(raw["quantity"])
        where = f"{sheet} row {row_no}"

        if not item_id:
            if title and (quantity or 0) >= 1:
                result.errors.append(f"{where}: {title!r} has no ID. Run scripts/assign_ids.py first.")
            elif not title and quantity:
                total = quantity
            continue
        if not id_re.match(item_id):
            result.errors.append(f"{where}: invalid ID {item_id!r}, expected {spec['prefix']}-0000")
            continue
        if item_id in seen_ids:
            result.errors.append(f"{where}: ID {item_id} is also used on {seen_ids[item_id]}")
            continue
        seen_ids[item_id] = where
        if not title:
            result.errors.append(f"{where}: {item_id} has no title")
            continue
        if quantity is None:
            result.warnings.append(f"{where}: {item_id} has no SUM value, imported with quantity 1")
            quantity = 1
        elif quantity == 0:
            result.warnings.append(f"{where}: {item_id} has SUM 0, imported with quantity 0")

        if "era" in spec:
            era = spec["era"]
        elif spec.get("era_from_title"):
            era = eras.from_title(title)
        else:
            era_raw = clean(raw.get("era"))
            if not era_raw:
                era = eras.fallback
                result.warnings.append(f"{where}: {item_id} {title!r} has no era, imported under {era!r}")
            else:
                era = eras.resolve(era_raw)
                if era is None:
                    unknown_eras[era_raw].append(row_no)
                    continue

        year, year_text, year_uncertain = parse_year(raw.get("year"))
        text = {f: clean(raw.get(f)) for f in TEXT_FIELDS}
        notes = [n for i, h in trailing if not h and (n := clean(row[i] if i < len(row) else None))]
        text["description"] = "; ".join(
            filter(None, [text["description"], clean(raw.get("additional_info")), *notes,
                          *(text[f] for f in spec.get("append_to_description", []))])) or None
        extra = {h: v for i, h in trailing if h and (v := clean(row[i] if i < len(row) else None))}

        item = {
            "id": item_id,
            "era": era,
            "format": match_format(spec.get("format_rules", []), spec["format"],
                                   text["description"]),
            "source_sheet": sheet,
            "source_row": row_no,
            "title": title,
            "year": year,
            "year_text": year_text,
            "year_uncertain": year_uncertain,
            **text,
            "catalog_norm": catalog_norm(text["catalog_no"]),
            "barcode": extract_barcode(text["catalog_no"], text["description"]),
            "quantity": quantity,
            "extra": extra,
        }
        # source_row is left out of the hash: moving a row in Excel is not a change.
        item["content_hash"] = content_hash({k: v for k, v in item.items() if k != "source_row"})
        result.items.append(item)
        count += quantity

    for name, rows in unknown_eras.items():
        rows_text = ", ".join(map(str, rows[:10])) + (" ..." if len(rows) > 10 else "")
        result.errors.append(f"{sheet}: unknown era {name!r} (rows {rows_text}). Add it to config/eras.yaml.")
    result.sheet_counts[sheet] = (count, total)
    if total is not None and total != count:
        result.warnings.append(f"{sheet}: imported quantity {count} differs from the SUM total row ({total})")


def read_concerts(sheet_rows, result):
    header_index = next((i for i, r in enumerate(sheet_rows) if "Day" in [clean(v) for v in r]), None)
    if header_index is None:
        result.errors.append("Concerts: header row with 'Day' not found")
        return
    header = [clean(v) for v in sheet_rows[header_index]]
    col = {name: header.index(name) for name in ["Day", "Month", "Year", "Place", "Venue", "Country", "Kees", "With whom"]
           if name in header}
    if len(col) < 8:
        result.errors.append(f"Concerts: expected columns Day, Month, Year, Place, Venue, Country, Kees, With whom; found {header}")
        return
    total = None
    for row_no, row in enumerate(sheet_rows[header_index + 1:], start=header_index + 2):
        if any(clean(v) == "Concerts in total" for v in row):
            total = parse_quantity(row[0])
            continue
        year = parse_quantity(row[col["Year"]])
        if year is None:
            continue
        try:
            day = date(year, parse_quantity(row[col["Month"]]), parse_quantity(row[col["Day"]]))
        except (TypeError, ValueError):
            result.errors.append(f"Concerts row {row_no}: invalid date")
            continue
        result.concerts.append({
            "date": day,
            "place": clean(row[col["Place"]]),
            "venue": clean(row[col["Venue"]]),
            "country": clean(row[col["Country"]]),
            "with_kees": parse_quantity(row[col["Kees"]]) == 1,
            "with_whom": clean(row[col["With whom"]]),
        })
    result.sheet_counts["Concerts"] = (len(result.concerts), total)
    if total is not None and total != len(result.concerts):
        result.warnings.append(f"Concerts: read {len(result.concerts)} concerts, total row says {total}")


def read(path, sheets_config, eras):
    workbook = load_workbook(path)
    result = Result()
    seen_ids = {}
    for spec in sheets_config["item_sheets"]:
        if spec["sheet"] not in workbook:
            result.errors.append(f"Sheet {spec['sheet']!r} not found in {Path(path).name}")
            continue
        read_items(workbook[spec["sheet"]], spec, eras, seen_ids, result)
    concerts_sheet = sheets_config["concerts"]["sheet"]
    if concerts_sheet in workbook:
        read_concerts(workbook[concerts_sheet], result)
    else:
        result.errors.append(f"Sheet {concerts_sheet!r} not found in {Path(path).name}")
    return result
